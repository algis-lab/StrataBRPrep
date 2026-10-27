#!/usr/bin/env python3
"""
Build index.html for the Geology Board Review site.

Reads the question bank spreadsheet, packs the questions into JSON,
gzips and encrypts them with the access PIN (AES-256-GCM, key from
PBKDF2-SHA256), and writes the result into template.html -> index.html.

Usage:
    python3 build.py "PRC_Geology_Board_Exam_Question_Bank_1.xlsx" [PIN]

Needs:  pip install openpyxl cryptography
"""
import base64, gzip, json, os, re, sys

import openpyxl
from cryptography.hazmat.primitives import hashes
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC

ITERATIONS = 400_000
HERE = os.path.dirname(os.path.abspath(__file__))
YEARS = ["2005", "2007", "2013", "2016", "2017", "2019", "2021", "2024", "2025"]
BASIS = {
    "Actual 2019 board paper (TQDS)": "A",
    "Examinee recall – reconstructed": "R",
    "Trend topic – constructed item": "T",
}


def clean(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Question Bank"]
    header = [clean(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    col = {h: i for i, h in enumerate(header)}
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[col["No."]] or not r[col["Question"]]:
            continue
        choices = [clean(r[col["Choice " + k]]) for k in "ABCD"]
        letter = clean(r[col["Answer"]]).upper()
        if letter not in "ABCD" or not letter:
            # fall back to matching the correct-answer text
            ans_text = clean(r[col["Correct Answer"]])
            letter = "ABCD"[choices.index(ans_text)]
        years = [y for y in YEARS if y in col and clean(r[col[y]])]
        out.append({
            "id": int(r[col["No."]]),
            "d": int(clean(r[col["Day"]]).split()[-1]),
            "s": clean(r[col["Subject"]]),
            "q": clean(r[col["Question"]]),
            "c": choices,
            "a": "ABCD".index(letter),
            "e": clean(r[col["Explanation"]]),
            "y": years,
            "b": BASIS.get(clean(r[col["Item Basis"]]), "R"),
            "n": clean(r[col["Notes / To Verify"]]),
        })
    return out


def encrypt(data: bytes, pin: str):
    salt, iv = os.urandom(16), os.urandom(12)
    kdf = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=salt, iterations=ITERATIONS)
    key = kdf.derive(pin.encode())
    ct = AESGCM(key).encrypt(iv, data, None)  # ciphertext || 16-byte tag (WebCrypto layout)
    b64 = lambda b: base64.b64encode(b).decode()
    return {"salt": b64(salt), "iv": b64(iv), "it": ITERATIONS, "ct": b64(ct)}


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    pin = sys.argv[2] if len(sys.argv) > 2 else "1212"
    qs = load(sys.argv[1])
    raw = json.dumps(qs, ensure_ascii=False, separators=(",", ":")).encode()
    blob = encrypt(gzip.compress(raw, 9), pin)
    tpl = open(os.path.join(HERE, "template.html"), encoding="utf-8").read()
    html = tpl.replace("/*__VAULT__*/null", json.dumps(blob))
    open(os.path.join(HERE, "index.html"), "w", encoding="utf-8").write(html)
    print(f"{len(qs)} questions -> index.html ({len(html)//1024} KB, PIN {'*'*len(pin)})")


if __name__ == "__main__":
    main()
